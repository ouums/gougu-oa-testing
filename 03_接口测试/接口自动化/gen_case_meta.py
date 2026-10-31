# -*- coding: utf-8 -*-
"""从用例表重新生成 case_meta.py（用例编号 -> Allure 严重级别）。

用法（在 03_接口测试/接口自动化/ 下执行）：

    python gen_case_meta.py

【为什么要生成而不是手写】
  严重级别（Allure 报告里每条用例左侧那个标签）的唯一权威来源是
  `02_测试用例/测试用例.xlsx` 的「优先级」列。手抄 135 条到代码里，
  下一次用例表调整就会静默对不上 —— 报告写 P1、用例表写「中」，
  而且没有任何东西会报错。生成出来的表可以随时重跑对齐。

【映射口径：三档对三档】
  高 -> critical (P1) / 中 -> normal (P2) / 低 -> minor (P3)
  **不启用 blocker(P0)** —— 用例表只有三档优先级，就一一对应三档；
  为了"看起来更全"把「高」升格成 P0 是没有依据的（本项目的定级口径
  A/B=严重、C/D=一般、E=轻微 是缺陷轴的，不是用例优先级轴的，两者不能混用）。
"""
import collections
import pathlib
import sys

HERE = pathlib.Path(__file__).resolve().parent
XLSX = HERE.parent.parent / "02_测试用例" / "测试用例.xlsx"
TARGET = HERE / "case_meta.py"

# 用例表「优先级」列的三个取值 -> Allure severity
LEVELS = {"高": "critical", "中": "normal", "低": "minor"}

HEADER = '''# -*- coding: utf-8 -*-
"""用例元数据（Allure 严重级别）—— **本文件由 gen_case_meta.py 生成，请勿手改**。

生成命令（在 03_接口测试/接口自动化/ 下执行）：

    python gen_case_meta.py

来源：`02_测试用例/测试用例.xlsx` 的「用例编号」列 -> 「优先级」列，共 %d 条（%s）。

映射：高 -> critical(P1) / 中 -> normal(P2) / 低 -> minor(P3)。
Allure 的 blocker(P0) **刻意不启用** —— 用例表只有三档优先级，就一一对应三档，
不为了"看起来更全"虚构第四档（理由详见 gen_case_meta.py 的模块 docstring）。

未被本表覆盖的用例不会带 severity 标签（Allure 按默认档显示），**不猜**。
"""

SEVERITY = {
'''


def main():
    try:
        import openpyxl
    except ImportError:
        sys.exit("缺少 openpyxl：python -m pip install openpyxl")

    if not XLSX.exists():
        sys.exit("找不到用例表：%s" % XLSX)

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]]

    def col(name):
        for i, v in enumerate(header):
            if v and name in str(v):
                return i
        sys.exit("用例表表头里找不到「%s」列，实测表头：%s" % (name, header))

    i_id, i_pr = col("编号"), col("优先级")

    sev, unknown = {}, collections.Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid, pr = row[i_id], row[i_pr]
        if not cid or not pr:
            continue
        cid, pr = str(cid).strip(), str(pr).strip()
        if pr not in LEVELS:
            unknown[pr] += 1
            continue
        sev[cid] = LEVELS[pr]

    if unknown:
        sys.exit("用例表出现未登记的优先级取值 %r —— 先确认口径再生成" % dict(unknown))
    if not sev:
        sys.exit("没有从用例表读到任何 (用例编号, 优先级) 对，请检查列名与表头")

    # 保持用例编号顺序（QX -> SP -> HT -> XZ 即用例表的自然顺序），便于人工核对
    lines = ["    %-10s %r,\n" % ('"%s":' % cid, lvl) for cid, lvl in sev.items()]
    dist = collections.Counter(sev.values())
    summary = " / ".join("%s %d" % (k, dist[k]) for k in ("critical", "normal", "minor")
                         if dist[k])

    TARGET.write_text(HEADER % (len(sev), summary) + "".join(lines) + "}\n", encoding="utf-8")
    print("已生成 %s：%d 条（%s）" % (TARGET.name, len(sev), summary))


if __name__ == "__main__":
    main()
